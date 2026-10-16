"""
ScopePilot AI — Document Extraction Service
Handles DOCX, PDF, TXT, XLSX with section-aware extraction,
OCR detection, and graceful error handling.
"""
import os
import logging
from typing import List, Dict, Any, Optional, Tuple
from pathlib import Path

logger = logging.getLogger(__name__)


class ExtractionError(Exception):
    pass


def extract_docx(file_path: str) -> Tuple[List[Dict[str, Any]], int]:
    """
    Extract text sections from a DOCX file.
    Returns (sections, page_count_estimate).
    Each section: {heading, content, section_order, char_count}
    """
    try:
        from docx import Document
        doc = Document(file_path)
    except Exception as e:
        raise ExtractionError(f"Cannot open DOCX file: {e}")

    sections = []
    current_heading = None
    current_content = []
    order = 0

    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        text = para.text.strip()
        if not text:
            continue

        if style_name.startswith("Heading"):
            # Save previous section
            if current_content:
                content = "\n".join(current_content)
                sections.append({
                    "heading": current_heading,
                    "content": content,
                    "section_order": order,
                    "char_count": len(content),
                    "page_number": None,
                    "sheet_name": None,
                })
                order += 1
            current_heading = text
            current_content = []
        else:
            current_content.append(text)

    # Save last section
    if current_content:
        content = "\n".join(current_content)
        sections.append({
            "heading": current_heading or "Document Body",
            "content": content,
            "section_order": order,
            "char_count": len(content),
            "page_number": None,
            "sheet_name": None,
        })

    # Also extract from tables
    for table in doc.tables:
        table_text = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            table_text.append(" | ".join(row_cells))
        if table_text:
            sections.append({
                "heading": "Table Data",
                "content": "\n".join(table_text),
                "section_order": order + 1,
                "char_count": len("\n".join(table_text)),
                "page_number": None,
                "sheet_name": None,
            })

    est_pages = max(1, sum(s["char_count"] for s in sections) // 3000)
    return sections, est_pages


def extract_pdf(file_path: str) -> Tuple[List[Dict[str, Any]], int]:
    """
    Extract text from PDF using PyMuPDF.
    Returns (sections, page_count).
    """
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise ExtractionError("PyMuPDF not installed. Run: pip install PyMuPDF")

    try:
        doc = fitz.open(file_path)
    except Exception as e:
        raise ExtractionError(f"Cannot open PDF: {e}")

    sections = []
    total_text = 0

    for page_num, page in enumerate(doc, start=1):
        text = page.get_text("text")
        if not text.strip():
            # Might be a scanned page
            sections.append({
                "heading": f"Page {page_num} (Possibly Scanned)",
                "content": "[No text extracted — this page may be a scanned image. OCR required.]",
                "section_order": page_num - 1,
                "char_count": 0,
                "page_number": page_num,
                "sheet_name": None,
                "ocr_needed": True,
            })
        else:
            total_text += len(text)
            sections.append({
                "heading": f"Page {page_num}",
                "content": text.strip(),
                "section_order": page_num - 1,
                "char_count": len(text),
                "page_number": page_num,
                "sheet_name": None,
                "ocr_needed": False,
            })

    return sections, len(doc)


def extract_txt(file_input) -> Tuple[List[Dict[str, Any]], int]:
    """Extract text from a plain text file path or bytes."""
    if isinstance(file_input, (bytes, bytearray)):
        content = file_input.decode("utf-8", errors="replace")
    elif isinstance(file_input, (str, Path)):
        try:
            with open(file_input, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
        except Exception as e:
            raise ExtractionError(f"Cannot read TXT file: {e}")
    else:
        raise ExtractionError("Invalid input for TXT extraction")

    if not content.strip():
        raise ExtractionError("The text file appears to be empty.")

    # Split by double newlines as sections
    raw_sections = content.split("\n\n")
    sections = []
    for i, section in enumerate(raw_sections):
        section = section.strip()
        if section:
            sections.append({
                "heading": f"Section {i + 1}",
                "content": section,
                "section_order": i,
                "char_count": len(section),
                "page_number": None,
                "sheet_name": None,
            })

    est_pages = max(1, len(content) // 3000)
    return sections, est_pages


def extract_xlsx(file_path: str) -> Tuple[List[Dict[str, Any]], int]:
    """Extract data from all sheets of an XLSX file."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ExtractionError("openpyxl not installed. Run: pip install openpyxl")

    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        raise ExtractionError(f"Cannot open XLSX: {e}")

    sections = []
    for order, sheet_name in enumerate(wb.sheetnames):
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            row_text = " | ".join(cells)
            if any(c.strip() for c in cells):
                rows.append(row_text)
        if rows:
            content = "\n".join(rows)
            sections.append({
                "heading": f"Sheet: {sheet_name}",
                "content": content,
                "section_order": order,
                "char_count": len(content),
                "page_number": None,
                "sheet_name": sheet_name,
            })

    return sections, len(wb.sheetnames)


def extract_document(
    file_path: str,
    filename: str,
) -> Dict[str, Any]:
    """
    Main extraction entry point.
    Auto-detects file type and returns structured extraction result.
    """
    ext = Path(filename).suffix.lower()
    supported = {".docx", ".pdf", ".txt", ".xlsx"}

    if ext not in supported:
        raise ExtractionError(
            f"Unsupported file type: {ext}. Supported types: {', '.join(supported)}"
        )

    if not os.path.exists(file_path):
        raise ExtractionError(f"File not found: {file_path}")

    file_size = os.path.getsize(file_path)
    if file_size == 0:
        raise ExtractionError("The uploaded file is empty.")

    try:
        if ext == ".docx":
            sections, page_count = extract_docx(file_path)
            method = "python-docx"
        elif ext == ".pdf":
            sections, page_count = extract_pdf(file_path)
            method = "PyMuPDF"
        elif ext == ".txt":
            sections, page_count = extract_txt(file_path)
            method = "plaintext"
        elif ext == ".xlsx":
            sections, page_count = extract_xlsx(file_path)
            method = "openpyxl"
        else:
            raise ExtractionError(f"Unhandled extension: {ext}")
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"Extraction failed: {e}")

    total_chars = sum(s.get("char_count", 0) for s in sections)
    ocr_needed = any(s.get("ocr_needed", False) for s in sections)

    return {
        "sections": sections,
        "page_count": page_count,
        "extraction_method": method,
        "total_characters": total_chars,
        "section_count": len(sections),
        "ocr_needed": ocr_needed,
        "warnings": ["Some pages appear to be scanned images. OCR is needed for full extraction."] if ocr_needed else [],
    }


def extract_template_structure(file_path: str) -> Dict[str, Any]:
    """
    Extract the structural metadata from a master SOW DOCX template.
    Returns headings, placeholder patterns, tables, styles.
    """
    try:
        from docx import Document
        import re
        doc = Document(file_path)
    except Exception as e:
        raise ExtractionError(f"Cannot read template DOCX: {e}")

    headings = []
    paragraphs_meta = []
    placeholder_pattern = re.compile(r"\{\{([A-Z_]+)\}\}")
    found_placeholders = set()
    static_sections = []
    tables_info = []

    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name if para.style else "Normal"

        # Detect headings
        if style_name.startswith("Heading"):
            level = style_name.replace("Heading", "").strip()
            headings.append({
                "text": text,
                "level": int(level) if level.isdigit() else 1,
                "paragraph_index": i,
                "style": style_name,
            })

        # Detect placeholders
        matches = placeholder_pattern.findall(text)
        for match in matches:
            found_placeholders.add(match)

        paragraphs_meta.append({
            "index": i,
            "style": style_name,
            "text_preview": text[:100],
            "has_placeholder": bool(matches),
            "placeholders": matches,
        })

    # Detect tables
    for t_idx, table in enumerate(doc.tables):
        cell_texts = []
        for row in table.rows:
            for cell in row.cells:
                cell_texts.append(cell.text.strip())
        table_placeholders = placeholder_pattern.findall(" ".join(cell_texts))
        tables_info.append({
            "table_index": t_idx,
            "row_count": len(table.rows),
            "col_count": len(table.columns),
            "placeholders": list(set(table_placeholders)),
            "preview": " | ".join(cell_texts[:6]),
        })

    # Detect styles used
    styles_used = list(set(
        p.style.name for p in doc.paragraphs if p.style
    ))

    return {
        "headings": headings,
        "placeholders": sorted(found_placeholders),
        "tables": tables_info,
        "styles_used": styles_used,
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "has_header": bool(doc.sections[0].header.paragraphs) if doc.sections else False,
        "has_footer": bool(doc.sections[0].footer.paragraphs) if doc.sections else False,
    }
