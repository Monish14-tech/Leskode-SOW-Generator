import os
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

from docx import Document
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)

# 1. Login
auth_res = client.post("/api/auth/login", json={"email": "admin@scopepilot.ai", "password": "password123"})
assert auth_res.status_code == 200, f"Login failed: {auth_res.text}"
token = auth_res.json()["access_token"]
headers = {"Authorization": f"Bearer {token}"}
print("1. Authenticated as admin@scopepilot.ai")

# 2. Create Project
proj_res = client.post("/api/projects", json={
    "client_name": "Quantum Logistics Group",
    "project_name": "Zoho One Global Supply Chain Suite",
    "description": "End-to-end CRM, automated inventory management, and financial billing system."
}, headers=headers)
assert proj_res.status_code == 201, f"Project creation failed: {proj_res.text}"
proj_id = proj_res.json()["id"]
print(f"2. Project created successfully: {proj_id}")

# 3. Upload Requirement Document
sample_text = """QUANTUM LOGISTICS GROUP — ZOHO IMPLEMENTATION SPECIFICATION
1. Business Background:
Quantum Logistics operates freight forwarding and cold-chain warehousing across 12 facilities.
Current operations rely on manual Excel workbooks and legacy paper billing causing fulfillment delays.

2. Sales & Customer Intake Scope:
• Capture freight inquiries from website form directly into Zoho CRM.
• Automatically assign deals to freight coordinators based on destination port within 30 minutes.
• If inquiry response exceeds 2 hours, alert General Manager.
• Prevent duplicate customer records by validating Company Tax ID and email.

3. Inventory & Warehouse Management:
• Multi-location warehouse stock tracking in Zoho Inventory.
• Automated purchase order creation when cold-storage consumables fall below threshold.

4. Financial Billing & Accounting:
• Automated GST tax invoicing in Zoho Books upon bill of lading generation.
• Bank reconciliation and automated overdue payment reminders via email and SMS.

5. Support & Claims:
• Freight loss and damage claim ticket tracking in Zoho Desk.
• Initial claim response SLA within 2 hours.

6. Integration & Migration:
• Webhook synchronization with Port Terminal EDI API via Zoho Flow.
• Historical migration of 25,000 shipment and customer records from legacy SQL database."""

files = {"file": ("quantum_requirements.txt", sample_text.encode("utf-8"), "text/plain")}
upload_res = client.post(f"/api/documents/{proj_id}/upload", files=files, headers=headers)
assert upload_res.status_code == 201, f"Upload failed: {upload_res.text}"
doc_id = upload_res.json()["id"]
print(f"3. Document uploaded: {doc_id}, status: {upload_res.json()['status']}")

# 4. Trigger AI Requirement Analysis
req_ana_res = client.post(f"/api/requirements/{proj_id}/analyse", json={}, headers=headers)
print("4. Requirements analysis triggered:", req_ana_res.json())

# Check requirements
reqs_res = client.get(f"/api/requirements/{proj_id}", headers=headers)
reqs_data = reqs_res.json()
reqs = reqs_data.get("items") if isinstance(reqs_data, dict) and "items" in reqs_data else (reqs_data.get("requirements", []) if isinstance(reqs_data, dict) else reqs_data)
print(f"   Extracted {len(reqs)} requirements:")
for r in reqs[:4]:
    print(f"   - [{r['req_id']}] {r['title']} ({r.get('category')}) | AI Rec: {r.get('is_ai_generated')}")
assert len(reqs) >= 3, "Should have extracted at least 3 requirements"

# 5. Trigger Module Recommendations
mod_rec_res = client.post(f"/api/modules/{proj_id}/recommend", headers=headers)
print("5. Module recommendation triggered:", mod_rec_res.json())

mods_res = client.get(f"/api/modules/{proj_id}", headers=headers)
mods_data = mods_res.json()
mods = mods_data if isinstance(mods_data, list) else (mods_data.get("modules") or mods_data.get("items", []))
print(f"   Recommended {len(mods)} Zoho modules:")
for m in mods:
    print(f"   - {m.get('module_name')}: {m.get('business_purpose', '')[:60]}... (Approved: {m.get('is_approved')})")
assert len(mods) >= 2, "Should have recommended at least 2 Zoho modules"

# Approve all modules
for m in mods:
    client.patch(f"/api/modules/{proj_id}/{m['id']}", json={"is_approved": True}, headers=headers)
print("   Approved all recommended modules for SOW.")

# 6. Generate Tasks for first module
task_res = client.post(f"/api/tasks/{proj_id}/generate", json={"module_mapping_id": mods[0]["id"]}, headers=headers)
print("6. Tasks generated for", mods[0].get("module_name"), ":", task_res.json())
tasks_res = client.get(f"/api/tasks/{proj_id}", headers=headers)
tasks_data = tasks_res.json()
tasks = tasks_data if isinstance(tasks_data, list) else (tasks_data.get("tasks") or tasks_data.get("items", []))
print(f"   Total scope tasks: {len(tasks)}")
assert len(tasks) >= 3, "Should have generated scope tasks"

# 7. Generate Automation Matrix
auto_res = client.post(f"/api/automations/{proj_id}/generate", headers=headers)
print("7. Automations generated:", auto_res.json())
autos_res = client.get(f"/api/automations/{proj_id}", headers=headers)
autos_data = autos_res.json()
autos = autos_data if isinstance(autos_data, list) else (autos_data.get("automations") or autos_data.get("items", []))
print(f"   Total automations: {len(autos)}")
assert len(autos) >= 1, "Should have generated automations"

# 8. Generate Clarifications & Integrations
clq_res = client.post(f"/api/integrations/{proj_id}/clarifications/generate", headers=headers)
client.post(f"/api/integrations/{proj_id}/integrations/generate", headers=headers)
clqs_res = client.get(f"/api/integrations/{proj_id}/clarifications", headers=headers)
clqs_data = clqs_res.json()
clqs = clqs_data if isinstance(clqs_data, list) else (clqs_data.get("questions") or clqs_data.get("items", []))
print(f"8. Generated {len(clqs)} clarification questions.")
assert len(clqs) >= 1, "Should have generated clarifications"

# 9. Run Validation Audit
val_res = client.post(f"/api/validation/{proj_id}/run", headers=headers)
val = val_res.json()
print(f"9. Validation Audit: Total Reqs={val['total_requirements']}, Unmapped={val['unmapped_requirements']}, Level={val['overall_level']}")

# 10. Generate SOW Document (DOCX)
sow_res = client.post(f"/api/sow/{proj_id}/generate", json={}, headers=headers)
assert sow_res.status_code == 201, f"SOW generation failed: {sow_res.text}"
sow_data = sow_res.json()
print(f"10. SOW v{sow_data['version_number']} generated, status={sow_data['status']}, docx_available={sow_data['docx_available']}")

# 11. Download DOCX
doc_download = client.get(f"/api/sow/{proj_id}/download/{sow_data['id']}", headers=headers)
assert doc_download.status_code == 200, f"Download failed: {doc_download.status_code}"
file_bytes = len(doc_download.content)
print(f"11. Downloaded DOCX file successfully: {file_bytes} bytes")
assert file_bytes > 1000, f"File too small: {file_bytes} bytes"

# Verify DOCX content
import tempfile
with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf:
    tf.write(doc_download.content)
    tmp_name = tf.name

try:
    sow_doc = Document(tmp_name)
    all_p_text = "\n".join(p.text for p in sow_doc.paragraphs)
    all_tbl_text = "\n".join(" | ".join(c.text for c in r.cells) for t in sow_doc.tables for r in t.rows)
    all_text = all_p_text + "\n" + all_tbl_text

    assert "Quantum Logistics Group" in all_text, "Client name missing from SOW"
    assert "Zoho One Global Supply Chain Suite" in all_text, "Project name missing from SOW"
    print("12. DOCX paragraphs and tables successfully verified: Found Client and Project name inside actual Word document!")
finally:
    if os.path.exists(tmp_name):
        os.remove(tmp_name)

print("\n" + "="*70)
print("   ALL 12 END-TO-END VERIFICATION STEPS PASSED WITH 100% SUCCESS!")
print("="*70)
