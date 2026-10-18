import os
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

from docx import Document
from fastapi.testclient import TestClient
from app.main import app

def test_monish_enterprises_sow_flow():
    client = TestClient(app)

    print("\n--- 1. Authenticate ---")
    auth_res = client.post("/api/auth/login", json={"email": "admin@scopepilot.ai", "password": "password123"})
    if auth_res.status_code != 200:
        auth_res = client.post("/api/auth/register", json={
            "email": "admin@scopepilot.ai",
            "password": "password123",
            "full_name": "Admin Consultant",
            "organization_name": "ScopePilot Enterprise"
        })
    assert auth_res.status_code == 200, f"Login/Register failed: {auth_res.text}"
    token = auth_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    print("   Authenticated as admin@scopepilot.ai")

    print("\n--- 2. Create Project: Monish Enterprises ---")
    proj_res = client.post("/api/projects", json={
        "client_name": "Monish Enterprises",
        "project_name": "Zoho Enterprise Operations Platform",
        "primary_zoho_edition": "Zoho One",
        "description": "End-to-end industrial steel trading, daily market price board, multi-godown stock, weighbridge gate entry, and GST invoicing."
    }, headers=headers)
    assert proj_res.status_code == 201, f"Project creation failed: {proj_res.text}"
    proj_id = proj_res.json()["id"]
    print(f"   Project created: {proj_id}")

    print("\n--- 3. Upload Requirements Document ---")
    monish_spec = """MONISH ENTERPRISES — ZOHO ONE IMPLEMENTATION SPECIFICATION
1. Executive Summary & Business Background:
Monish Enterprises is an industrial distributor and dealer of steel products, pipes, roofing sheets, fasteners, and hardware components based in Coimbatore, Tamil Nadu.
Currently, operations are handled through manual registers, offline spreadsheets, and disconnected phone/WhatsApp messages, leading to quote delays, inventory discrepancies, and lack of real-time margin visibility.
The objective is to deploy an end-to-end connected Zoho ecosystem connecting customer inquiry through pricing, quotation, inventory allocation, dispatch, GST invoicing, and executive margin dashboards.

2. Lead Intake & Sales Pipeline (Zoho CRM):
• Centralize inbound inquiries from IndiaMART, website contact forms, WhatsApp Business, phone calls, and walk-in trade customers into Zoho CRM Leads.
• Automated round-robin lead assignment to sales coordinators with deduplication on GST number, mobile, and email.
• Follow-up activity scheduling, call logs, and SLA alert to General Manager if a lead is untouched for more than 45 minutes.
• Deal pipeline stages: Lead Qualified -> Site / Requirement Analysis -> Quotation Submitted -> Price Negotiation -> Order Confirmed -> Closed Won / Lost.
• Order Confirmed status triggers automated Sales Order creation in Zoho Books and cross-team project creation in Zoho Projects.

3. Quotation & Daily Price Engine (Zoho Creator & CRM):
• In the steel and industrial hardware trade, purchase rates change daily from mills and manufacturers.
• Daily Market Price Board built in Zoho Creator where daily mill rates, landed freight costs, and category margins (Retail, Builder, Fabricator, Wholesale Dealer) are set each morning and approved by the Owner.
• CRM Quotations automatically pull today's approved rate for the selected customer price category.
• Margin Guard: Any quotation discount below the configured margin floor requires mobile approval from the Owner.

4. Financial Accounting & Invoicing (Zoho Books):
• Generation of statutory GST-compliant invoices with automatic HSN codes, e-invoice (IRN), and e-way bill from confirmed Sales Orders.
• Vendor Purchase Orders linked to mill accounts, three-way matching (PO -> Goods Receipt Note -> Vendor Bill) before payment release.
• Customer credit limits, payment terms, and automated payment reminder notifications for overdue accounts.
• Automated bank feed reconciliation and monthly statutory GST return summaries.

5. Inventory & Multi-Godown Stock Management (Zoho Inventory):
• Centralized item master organized into item groups (Structural Steel, TMT Bars, Roofing Sheets, Bolts & Fasteners, Industrial Tools).
• Multi-unit handling with automatic unit conversion (metric tonnes, kilograms, bundles, pieces, square feet).
• Multi-godown tracking across Main Central Yard, Processing Warehouse, and City Trade Branch with inter-godown transfer orders.
• Heat number, batch number, and mill test certificate (MTC) tracking upon stock inwarding.
• Automated reorder alerts when inventory levels breach minimum safety stock thresholds.

6. Fulfillment, Weighbridge & Gate Management (Zoho Creator & Inventory):
• Fulfillment orders created upon Deal confirmation in CRM.
• Vehicle Inward Gate Entry referencing Purchase Order and capturing weighbridge slip gross/tare weight.
• Vehicle Outward Gate Pass referencing Sales Order with loaded weight vs billed weight verification. Dispatches exceeding tolerance held for supervisor approval.
• Accepted GRN quantities update Zoho Inventory physical stock in real time.

7. Human Resources & Payroll (Zoho People & Zoho Payroll):
• Employee records, department hierarchy, shift schedules, and geofenced attendance tracking for yard and field sales staff in Zoho People.
• Leave management with approval workflows (Employee -> Manager -> HR).
• Monthly attendance and LOP data synchronized to Zoho Payroll for automated salary calculation, statutory deductions (PF, ESI, PT), and payslip generation.
• Completed pay run posts journal entries directly to Zoho Books.

8. Cross-Module Collaboration & Executive Intelligence (Zoho Projects & Analytics):
• Automated Zoho Projects workspace created for confirmed high-value and custom fabrication orders.
• Executive Dashboards in Zoho Analytics: Daily Realized Margin vs Replacement Margin, Sales Rep Conversion Ratios, Stock Ageing & Revaluation Report, and Customer Outstanding Ageing.

9. Project Assumptions & Exclusions:
• Client will provide single Point of Contact (SPOC) and master data in CSV format.
• Zoho licenses and third-party WhatsApp Business API fees are procured directly by Monish Enterprises.
• Weighbridge hardware and barcode printer procurement is outside software scope."""

    files = {"file": ("monish_enterprises_requirements.txt", monish_spec.encode("utf-8"), "text/plain")}
    up_res = client.post(f"/api/documents/{proj_id}/upload", files=files, headers=headers)
    assert up_res.status_code == 201, f"Upload failed: {up_res.text}"
    doc_id = up_res.json()["id"]
    print(f"   Document uploaded successfully: {doc_id}")

    print("\n--- 4. Extract Requirements via Gemini AI ---")
    ana_res = client.post(f"/api/requirements/{proj_id}/analyse", headers=headers)
    assert ana_res.status_code == 200, f"Analysis failed: {ana_res.text}"

    req_list = client.get(f"/api/requirements/{proj_id}", headers=headers)
    reqs = req_list.json().get("items", []) or req_list.json().get("requirements", [])
    print(f"   Requirements extracted: {len(reqs)}")
    assert len(reqs) >= 8, f"Expected at least 8 requirements, got {len(reqs)}"
    for r in reqs[:4]:
        print(f"   - [{r['req_id']}] {r['title']} ({r['category']}) | AI Rec: {r['is_ai_generated']}")

    print("\n--- 5. Module Recommendations via Gemini AI ---")
    mod_res = client.post(f"/api/modules/{proj_id}/recommend", headers=headers)
    assert mod_res.status_code == 200, f"Module recommendation failed: {mod_res.text}"

    get_mods = client.get(f"/api/modules/{proj_id}", headers=headers)
    modules = get_mods.json() if isinstance(get_mods.json(), list) else get_mods.json().get("modules", [])
    print(f"   Recommended {len(modules)} Zoho modules:")
    for m in modules:
        print(f"   - {m['module_name']}: {m['business_purpose'][:65]}...")
        # Approve all modules
        client.patch(f"/api/modules/{proj_id}/{m['id']}", json={"is_approved": True}, headers=headers)
    print("   Approved all recommended modules for SOW.")

    print("\n--- 6. Generate Scope Tasks for Modules ---")
    for m in modules[:3]:
        t_res = client.post(f"/api/tasks/{proj_id}/generate", json={"module_mapping_id": m["id"]}, headers=headers)
        assert t_res.status_code == 200
    all_tasks = client.get(f"/api/tasks/{proj_id}", headers=headers)
    tasks = all_tasks.json() if isinstance(all_tasks.json(), list) else all_tasks.json().get("tasks", [])
    print(f"   Total scope tasks generated: {len(tasks)}")

    print("\n--- 7. Generate Automation Matrix ---")
    auto_res = client.post(f"/api/automations/{proj_id}/generate", headers=headers)
    assert auto_res.status_code == 200
    all_autos = client.get(f"/api/automations/{proj_id}", headers=headers)
    autos = all_autos.json() if isinstance(all_autos.json(), list) else all_autos.json().get("automations", [])
    print(f"   Total automations generated: {len(autos)}")

    print("\n--- 8. Generate Clarifications ---")
    clq_res = client.post(f"/api/integrations/{proj_id}/clarifications/generate", headers=headers)
    assert clq_res.status_code == 200
    all_clq = client.get(f"/api/integrations/{proj_id}/clarifications", headers=headers)
    clqs = all_clq.json().get("questions", []) if isinstance(all_clq.json(), dict) else all_clq.json()
    print(f"   Generated {len(clqs)} clarification questions.")

    print("\n--- 9. Run Scope Coverage Validation ---")
    val_res = client.post(f"/api/validation/{proj_id}/run", headers=headers)
    assert val_res.status_code == 200
    val_data = val_res.json()
    print(f"   Validation Level: {val_data['overall_level']}, Total Reqs: {val_data['total_requirements']}")

    print("\n--- 10. Generate SOW Document (DOCX) using Leskode Master Template ---")
    sow_res = client.post(f"/api/sow/{proj_id}/generate", headers=headers)
    assert sow_res.status_code == 201, f"SOW generation failed: {sow_res.text}"
    sow_data = sow_res.json()
    sow_id = sow_data["id"]
    print(f"   SOW generated successfully: ID={sow_id}, version={sow_data['version_number']}, docx_available={sow_data['docx_available']}")

    print("\n--- 11. Download DOCX ---")
    dl_res = client.get(f"/api/sow/{proj_id}/download/{sow_id}", headers=headers)
    assert dl_res.status_code == 200, f"Download failed: {dl_res.status_code}"
    file_bytes = len(dl_res.content)
    print(f"   Downloaded SOW DOCX file: {file_bytes} bytes")
    assert file_bytes > 20000, "DOCX file is suspiciously small"

    # Save to disk for inspection
    out_file = Path("monish_enterprises_sow_verified.docx")
    out_file.write_bytes(dl_res.content)

    print("\n--- 12. Inspect DOCX Contents & Master Template Compliance ---")
    doc = Document("monish_enterprises_sow_verified.docx")
    full_text = "\n".join([p.text for p in doc.paragraphs])
    table_text = "\n".join([" | ".join([c.text.strip() for c in row.cells]) for t in doc.tables for row in t.rows])
    combined = full_text + "\n" + table_text

    print(f"   Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")

    # Verification checks
    assert "Monish Enterprises" in combined, "Missing client name 'Monish Enterprises'"
    assert "SCOPE OF WORK" in combined, "Missing 'SCOPE OF WORK'"
    assert "Leskode" in combined, "Missing 'Leskode'"
    assert "Shavs Technologies Pvt Ltd" in combined, "Missing Shavs Technologies Pvt Ltd"
    assert "Confidentiality Notice" in combined, "Missing Confidentiality Notice"
    assert "Vaishnavi T" in combined, "Missing Vaishnavi T contact block"
    assert "+91 9500424307" in combined, "Missing contact number +91 9500424307"
    assert "Executive Summary" in combined, "Missing Executive Summary"
    assert "Process Flow" in combined, "Missing Process Flow"
    assert "Acceptance & Sign-Off" in combined, "Missing Acceptance & Sign-Off"
    assert "Data migration, hardware, and post-go-live support beyond hypercare" in combined, "Missing locked exclusions closing line"

    print("\n>>> ALL 12 CHECKS PASSED FOR MONISH ENTERPRISES SOW GENERATION! <<<")

if __name__ == "__main__":
    test_monish_enterprises_sow_flow()
